"""Socle des rendus tableur : un **tableau** composé une fois, rendu en CSV ou en xlsx (E16US016).

Pourquoi un socle et non un adapter par document et par format : ADR-0101 §6.

⚠️ La neutralisation des formules (CWE-1236) n'est **pas** commune aux deux rendus : le CSV préfixe
une apostrophe, le xlsx force le type de la cellule. Appliquer le remède CSV au xlsx afficherait
l'apostrophe dans la case, sous les yeux de l'organisateur.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from typing import Protocol

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from infrastructure.erreurs import InfrastructureError

_SEPARATEUR = ";"

# Caractères par lesquels un tableur reconnaît une **formule**. Une cellule qui commence par l'un
# d'eux est exécutée à l'ouverture (CWE-1236) — or les noms d'archers et de clubs viennent de la
# saisie et de l'import FFTA, et le fichier est fait pour partir chez la trésorière.
_AMORCES_DE_FORMULE = ("=", "+", "-", "@", "\t", "\r")

_FORMAT_MONTANT_XLSX = "0.00"


@dataclass(frozen=True)
class Montant:
    """Une somme, en centimes. Marque la cellule comme **nombre**, jamais comme texte.

    ⚠️ C'est ce marquage, et lui seul, qui rend la colonne sommable : un montant suffixé ou
    neutralisé redevient du texte, ce qu'ADR-0101 §4 interdit — l'usage de l'export *est* la somme.
    """

    centimes: int


Cellule = str | int | Montant
"""Ce qu'une case peut porter. `str` est du texte (donc neutralisable), le reste est numérique."""


@dataclass(frozen=True)
class Tableau:
    """Un document tabulaire prêt à rendre : une ligne d'en-tête, puis les données.

    ⚠️ Aucun total, aucun bloc, aucun titre de document (ADR-0101 §4) : ce qui fait un beau PDF
    casse le tri et le filtre d'un tableur.
    """

    entetes: tuple[str, ...]
    lignes: tuple[tuple[Cellule, ...], ...]


class RenduTableur(Protocol):
    """Port interne : transforme un `Tableau` en octets téléchargeables."""

    def __call__(self, tableau: Tableau) -> bytes:
        """Rend le tableau ; lève `InfrastructureError` si l'écriture échoue."""
        ...


def rendre_csv(tableau: Tableau) -> bytes:
    """Rend en CSV lisible par un tableur français (ADR-0101 §4 : BOM, point-virgule, virgule)."""
    try:
        tampon = io.StringIO(newline="")
        # `\r\n` : fin de ligne attendue par la RFC 4180 et par Excel sous Windows.
        redacteur = csv.writer(tampon, delimiter=_SEPARATEUR, lineterminator="\r\n")
        redacteur.writerow(tableau.entetes)
        redacteur.writerows(
            tuple(_cellule_csv(cellule) for cellule in ligne) for ligne in tableau.lignes
        )
        return tampon.getvalue().encode("utf-8-sig")
    except Exception as echec:  # pragma: no cover - défense, `csv` n'échoue pas sur du `str`
        raise InfrastructureError("Échec du rendu CSV du document.") from echec


def rendre_xlsx(tableau: Tableau) -> bytes:
    """Rend en classeur Excel — une seule feuille, l'en-tête figé."""
    try:
        classeur = Workbook()
        feuille = classeur.active
        feuille.append(list(tableau.entetes))
        for numero, ligne in enumerate(tableau.lignes, start=2):
            _ecrire_ligne_xlsx(feuille, numero, ligne)
        # Fige l'en-tête : un journal de mille lignes se lit en défilant, sans perdre ses colonnes.
        feuille.freeze_panes = "A2"
        tampon = io.BytesIO()
        classeur.save(tampon)
        return tampon.getvalue()
    except Exception as echec:  # pragma: no cover - défense, openpyxl n'échoue pas sur ces types
        raise InfrastructureError("Échec du rendu xlsx du document.") from echec


def _ecrire_ligne_xlsx(feuille: Worksheet, numero: int, ligne: tuple[Cellule, ...]) -> None:
    """Écrit une ligne, en **forçant le type** de chaque case de texte.

    ⚠️ openpyxl interprète toute chaîne commençant par `=` comme une formule (`data_type` vaut
    alors `f`) : sans ce forçage, un nom d'archer importé exécuterait du calcul à l'ouverture.
    Le `data_type` se pose **après** l'affectation, qui le recalcule.
    """
    for colonne, valeur in enumerate(ligne, start=1):
        if isinstance(valeur, Montant):
            case = feuille.cell(row=numero, column=colonne, value=valeur.centimes / 100)
            case.number_format = _FORMAT_MONTANT_XLSX
            continue
        case = feuille.cell(row=numero, column=colonne, value=valeur)
        if isinstance(valeur, str):
            case.data_type = "s"


def _cellule_csv(cellule: Cellule) -> str:
    """Rend une case en texte CSV : montants à la française, texte neutralisé."""
    if isinstance(cellule, Montant):
        # Virgule décimale et **sans symbole** — pour rester sommable au tableur.
        return f"{cellule.centimes / 100:.2f}".replace(".", ",")
    if isinstance(cellule, int):
        return str(cellule)
    return f"'{cellule}" if cellule.startswith(_AMORCES_DE_FORMULE) else cellule
