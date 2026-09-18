"""Tests du rendu **xlsx** (E16US016) — écrits après l'adapter : il n'y a pas d'oracle en jeu.

Ce qui se vérifie ici est ce qu'aucun test de service ne peut voir : le fichier produit s'ouvre,
ses nombres sont des nombres, et **son texte n'est pas du calcul**. Ce dernier point est le motif
principal du fichier : openpyxl exécute par défaut toute chaîne commençant par `=` (CWE-1236), et
les noms d'archers viennent de l'import FFTA.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from infrastructure.tableur.tableau import Montant, Tableau, rendre_csv, rendre_xlsx


def _feuille(octets: bytes) -> Worksheet:
    classeur = load_workbook(io.BytesIO(octets))
    return classeur.active


def test_l_entete_et_les_lignes_sont_rendues() -> None:
    tableau = Tableau(("Nom", "Club"), (("MARTIN", "Kervignarc"), ("CADIOU", "Quimper")))

    feuille = _feuille(rendre_xlsx(tableau))

    assert [case.value for case in feuille["1"]] == ["Nom", "Club"]
    assert [case.value for case in feuille["2"]] == ["MARTIN", "Kervignarc"]
    assert feuille.max_row == 3


def test_un_nom_qui_commence_par_egal_reste_du_texte() -> None:
    """⚠️ **Le garde-fou de ce module** (CWE-1236, ADR-0101 §4 bis, versant tableur).

    Vérifié en réintroduisant le bug : sans le forçage de `data_type`, `data_type` vaut `f` et la
    case devient une formule que le tableur de la trésorière exécute à l'ouverture.
    """
    tableau = Tableau(("Club",), (("=1+1",), ("+33 2 98",), ("-Kervignarc",), ("@somme",)))

    feuille = _feuille(rendre_xlsx(tableau))

    for rang in (2, 3, 4, 5):
        case = feuille.cell(row=rang, column=1)
        assert case.data_type == "s", f"ligne {rang} : {case.value!r} est passée en formule"


def test_le_texte_neutralise_n_est_pas_defigure() -> None:
    """Contrepartie du test précédent : le xlsx **n'ajoute pas** l'apostrophe du CSV.

    ⚠️ C'est ce qui interdit de partager la neutralisation entre les deux rendus : appliquer le
    remède CSV ici afficherait `'=1+1` dans la case, sous les yeux de l'organisateur.
    """
    feuille = _feuille(rendre_xlsx(Tableau(("Club",), (("=1+1",),))))

    assert feuille.cell(row=2, column=1).value == "=1+1"


def test_un_montant_est_un_nombre_sommable() -> None:
    """L'usage de cet export *est* la somme (ADR-0101 §4) : un montant texte la rendrait fausse."""
    feuille = _feuille(rendre_xlsx(Tableau(("Dû",), ((Montant(850),), (Montant(-500),)))))

    assert feuille.cell(row=2, column=1).value == pytest.approx(8.50)
    assert feuille.cell(row=3, column=1).value == pytest.approx(-5.00)
    assert feuille.cell(row=2, column=1).data_type == "n"


def test_un_entier_reste_un_nombre() -> None:
    """Un numéro de départ ou de cible se trie en ordre numérique, pas alphabétique (2 avant 10)."""
    feuille = _feuille(rendre_xlsx(Tableau(("Cible",), ((10,), (2,)))))

    assert [feuille.cell(row=rang, column=1).value for rang in (2, 3)] == [10, 2]


def test_l_entete_est_fige() -> None:
    """Un journal de mille lignes se lit en défilant — sans ses colonnes, il ne se lit pas."""
    assert _feuille(rendre_xlsx(Tableau(("Nom",), (("MARTIN",),)))).freeze_panes == "A2"


def test_un_tableau_sans_ligne_produit_un_classeur_valide() -> None:
    """Un tournoi sans acte tracé rend un document **vide mais ouvrable**, pas un fichier cassé."""
    feuille = _feuille(rendre_xlsx(Tableau(("Horodatage", "Auteur"), ())))

    assert [case.value for case in feuille["1"]] == ["Horodatage", "Auteur"]
    assert feuille.max_row == 1


def test_les_deux_rendus_partent_du_meme_tableau() -> None:
    """Le format n'agit qu'au **rendu** : mêmes données, deux mises en forme (ADR-0101 §4).

    On ne compare pas les octets — ils n'ont aucune raison de se ressembler — mais le fait que le
    même `Tableau` alimente les deux — la propriété qu'une seule classe compositrice tient.
    """
    tableau = Tableau(("Nom", "Dû"), (("MARTIN", Montant(850)),))

    csv = rendre_csv(tableau).decode("utf-8-sig")
    feuille = _feuille(rendre_xlsx(tableau))

    assert "MARTIN;8,50" in csv
    assert feuille.cell(row=2, column=1).value == "MARTIN"
    assert feuille.cell(row=2, column=2).value == pytest.approx(8.50)
