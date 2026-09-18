"""Tests du rendu **xlsx** (E16US016) — écrits après l'adapter : il n'y a pas d'oracle en jeu.

Ce qui se vérifie ici est ce qu'aucun test de service ne peut voir : le fichier produit s'ouvre,
ses nombres sont des nombres, et **son texte n'est pas du calcul**. Ce dernier point est le motif
principal du fichier : openpyxl exécute par défaut toute chaîne commençant par `=` (CWE-1236), et
les noms d'archers viennent de l'import FFTA.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from infrastructure.tableur.grille import Grille, Montant, rendre_csv, rendre_xlsx


def _feuille(octets: bytes) -> Worksheet:
    classeur = load_workbook(io.BytesIO(octets))
    return classeur.active


def test_l_entete_et_les_lignes_sont_rendues() -> None:
    grille = Grille(("Nom", "Club"), (("MARTIN", "Kervignarc"), ("CADIOU", "Quimper")))

    feuille = _feuille(rendre_xlsx(grille))

    assert [case.value for case in feuille["1"]] == ["Nom", "Club"]
    assert [case.value for case in feuille["2"]] == ["MARTIN", "Kervignarc"]
    assert feuille.max_row == 3


def test_un_nom_qui_commence_par_egal_reste_du_texte() -> None:
    """⚠️ **Le garde-fou de ce module** (CWE-1236, ADR-0101 §4 bis, versant tableur).

    Vérifié en réintroduisant le bug : sans le forçage de `data_type`, `data_type` vaut `f` et la
    case devient une formule que le tableur de la trésorière exécute à l'ouverture.
    """
    grille = Grille(("Club",), (("=1+1",), ("+33 2 98",), ("-Kervignarc",), ("@somme",)))

    feuille = _feuille(rendre_xlsx(grille))

    for rang in (2, 3, 4, 5):
        case = feuille.cell(row=rang, column=1)
        assert case.data_type == "s", f"ligne {rang} : {case.value!r} est passée en formule"


def test_le_texte_neutralise_n_est_pas_defigure() -> None:
    """Contrepartie du test précédent : le xlsx **n'ajoute pas** l'apostrophe du CSV.

    ⚠️ C'est ce qui interdit de partager la neutralisation entre les deux rendus : appliquer le
    remède CSV ici afficherait `'=1+1` dans la case, sous les yeux de l'organisateur.
    """
    feuille = _feuille(rendre_xlsx(Grille(("Club",), (("=1+1",),))))

    assert feuille.cell(row=2, column=1).value == "=1+1"


def test_un_montant_est_un_nombre_sommable() -> None:
    """L'usage de cet export *est* la somme (ADR-0101 §4) : un montant texte la rendrait fausse."""
    feuille = _feuille(rendre_xlsx(Grille(("Dû",), ((Montant(850),), (Montant(-500),)))))

    assert feuille.cell(row=2, column=1).value == pytest.approx(8.50)
    assert feuille.cell(row=3, column=1).value == pytest.approx(-5.00)
    assert feuille.cell(row=2, column=1).data_type == "n"


def test_un_entier_reste_un_nombre() -> None:
    """Un numéro de départ ou de cible se trie en ordre numérique, pas alphabétique (2 avant 10)."""
    feuille = _feuille(rendre_xlsx(Grille(("Cible",), ((10,), (2,)))))

    assert [feuille.cell(row=rang, column=1).value for rang in (2, 3)] == [10, 2]


def test_l_entete_est_fige() -> None:
    """Un journal de mille lignes se lit en défilant — sans ses colonnes, il ne se lit pas."""
    assert _feuille(rendre_xlsx(Grille(("Nom",), (("MARTIN",),)))).freeze_panes == "A2"


def test_une_grille_sans_ligne_produit_un_classeur_valide() -> None:
    """Un tournoi sans acte tracé rend un document **vide mais ouvrable**, pas un fichier cassé."""
    feuille = _feuille(rendre_xlsx(Grille(("Horodatage", "Auteur"), ())))

    assert [case.value for case in feuille["1"]] == ["Horodatage", "Auteur"]
    assert feuille.max_row == 1


def test_les_deux_rendus_partent_de_la_meme_grille() -> None:
    """Le format n'agit qu'au **rendu** : mêmes données, deux mises en forme (ADR-0101 §4).

    On ne compare pas les octets — ils n'ont aucune raison de se ressembler — mais le fait que le
    même `Grille` alimente les deux — la propriété qu'une seule classe compositrice tient.
    """
    grille = Grille(("Nom", "Dû"), (("MARTIN", Montant(850)),))

    csv = rendre_csv(grille).decode("utf-8-sig")
    feuille = _feuille(rendre_xlsx(grille))

    assert "MARTIN;8,50" in csv
    assert feuille.cell(row=2, column=1).value == "MARTIN"
    assert feuille.cell(row=2, column=2).value == pytest.approx(8.50)


# --- E16US016, 2ᵉ passe de revue : l'en-tête est une ligne comme les autres ----------------------
#
# ⚠️ Ces trois cas ferment un trou que la 1ʳᵉ livraison laissait ouvert **dans les deux rendus** :
# l'en-tête court-circuitait la neutralisation, et l'axe adversarial l'a prouvé sur le XML produit
# (`<f>SUM(A1)</f>`, une formule réelle). Inoffensif tant que les en-têtes sont des constantes,
# exploitable au premier document à colonnes calculées — que le socle existe pour accueillir.


def test_un_entete_qui_commence_par_egal_reste_du_texte_en_xlsx() -> None:
    feuille = _feuille(rendre_xlsx(Grille(("=SUM(A1)", "Club"), ())))

    assert feuille.cell(row=1, column=1).data_type == "s"
    assert feuille.cell(row=1, column=1).value == "=SUM(A1)"


def test_un_entete_qui_commence_par_egal_est_neutralise_en_csv() -> None:
    entete = rendre_csv(Grille(("=SUM(A1)", "Club"), ())).decode("utf-8-sig").splitlines()[0]

    assert entete == "'=SUM(A1);Club"


def test_la_feuille_porte_le_titre_du_document() -> None:
    """La justification d'`openpyxl` (règle 11) promet « une feuille nommée » — elle l'est."""
    assert _feuille(rendre_xlsx(Grille(("Nom",), (), titre="Palmarès"))).title == "Palmarès"


def test_un_caractere_de_controle_est_retire_au_lieu_de_faire_tomber_l_export() -> None:
    """⚠️ openpyxl **refuse** ce que Python accepte dans une `str` (`IllegalCharacterError`).

    La 1ʳᵉ rédaction épinglait l'échec comme attendu : un seul caractère invisible dans un nom
    importé de la FFTA faisait tomber **tout** l'export xlsx du tournoi, quand le CSV du même
    document passait. On assainit plutôt — un onglet ou un nom nettoyé coûte moins qu'un export
    perdu le jour J (relevé en 2ᵉ passe, axes C1 et adversarial).
    """
    feuille = _feuille(rendre_xlsx(Grille(("Nom",), ((f"bon{chr(11)}jour",),))))

    assert feuille.cell(row=2, column=1).value == "bonjour"


def test_un_titre_interdit_par_excel_est_assaini_au_lieu_de_lever() -> None:
    """⚠️ Trois modes de panne **mesurés** par l'axe adversarial, tous fermés ici.

    Un `/` (libellé « Senior H / Arc classique ») levait `ValueError` → 500 sur tout l'export ;
    au-delà de 31 caractères openpyxl ne rend qu'un avertissement et Excel peut refuser le
    fichier ; un caractère de contrôle produisait un `workbook.xml` **mal formé**, donc un
    classeur illisible servi en 200 — pire que l'échec, parce qu'invisible.
    """
    assert (
        _feuille(rendre_xlsx(Grille(("Nom",), (), titre="Senior H / Arc"))).title
        == "Senior H   Arc"
    )
    assert len(_feuille(rendre_xlsx(Grille(("Nom",), (), titre="T" * 40))).title) == 31
    assert _feuille(rendre_xlsx(Grille(("Nom",), (), titre=f"a{chr(11)}b"))).title == "ab"
    assert _feuille(rendre_xlsx(Grille(("Nom",), (), titre="///"))).title == "Export"
